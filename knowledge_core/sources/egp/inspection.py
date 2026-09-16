from __future__ import annotations

import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from knowledge_core.sources.egp.models import CEFR_LEVELS, EGPCategoryConfig, EGPConfig
from knowledge_core.sources.egp.normalizer import normalize_empty
from knowledge_core.sources.egp.parser import (
    CAN_DO_ALIASES,
    FIELD_ALIASES,
    EGPParseError,
    _find_header_row,
    _resolve_field_columns,
    normalize_column_name,
)
from knowledge_core.sources.egp.paths import DEFAULT_EGP_REPORTS_DIR, DEFAULT_RAW_GRAMMAR_DIR, raw_file_path


SOURCE_INSPECTION_REPORT_NAME = "source_inspection_report.json"


def inspect_configured_sources(
    config: EGPConfig,
    raw_dir: str | Path = DEFAULT_RAW_GRAMMAR_DIR,
    *,
    category_id: str | None = None,
) -> dict[str, Any]:
    raw_root = Path(raw_dir)
    selected_categories = config.select_categories(category_id)
    configured_paths = {
        raw_file_path(raw_root, category).resolve(): category
        for category in selected_categories
    }
    files = [
        inspect_source_file(path, category)
        for path, category in sorted(configured_paths.items(), key=lambda item: str(item[0]))
    ]

    discovered_extra_files = []
    if raw_root.exists() and category_id is None:
        for path in sorted(raw_root.rglob("*.xlsx")):
            if path.resolve() not in configured_paths:
                discovered_extra_files.append(str(path))

    return {
        "source": config.source.name,
        "raw_dir": str(raw_root),
        "totals": _inspection_totals(files),
        "configured_categories": [category.id for category in selected_categories],
        "discovered_extra_files": discovered_extra_files,
        "files": files,
    }


def inspect_source_file(path: Path, category: EGPCategoryConfig) -> dict[str, Any]:
    base = {
        "category_id": category.id,
        "filename": path.name,
        "path": str(path),
        "exists": path.exists(),
        "file_size": path.stat().st_size if path.exists() else 0,
        "worksheet_names": [],
        "merged_ranges": [],
        "row_count": 0,
        "data_row_count": 0,
        "column_names": [],
        "super_category_values": {},
        "sub_category_values": {},
        "cefr_distribution": {},
        "duplicate_row_count": 0,
        "empty_rows": 0,
        "malformed_rows": [],
        "category_mismatch_rows": [],
        "issues": [],
    }
    if not path.exists():
        base["issues"].append({"code": "missing_file", "severity": "error"})
        return base
    if path.stat().st_size == 0:
        base["issues"].append({"code": "empty_file", "severity": "error"})
        return base

    try:
        workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception as exc:
        base["issues"].append(
            {
                "code": "malformed_workbook",
                "severity": "error",
                "message": str(exc),
            },
        )
        return base

    try:
        base["worksheet_names"] = list(workbook.sheetnames)
        worksheet = workbook.active
        base["merged_ranges"] = [str(cell_range) for cell_range in worksheet.merged_cells.ranges]
        rows = list(worksheet.iter_rows(values_only=True))
        base["row_count"] = len(rows)
        if not rows:
            base["issues"].append({"code": "empty_workbook", "severity": "error"})
            return base

        header_row_number, headers = _find_header_row(iter(rows))
        field_columns = _resolve_field_columns(headers)
        base["column_names"] = headers
        base["issues"].extend(_schema_issues(headers, field_columns))
        stats = _inspect_rows(
            rows[header_row_number:],
            headers,
            field_columns,
            category,
            header_row_number,
        )
        base.update(stats)
    except EGPParseError as exc:
        base["issues"].append(
            {"code": "unexpected_schema", "severity": "error", "message": str(exc)},
        )
    finally:
        workbook.close()
    return base


def write_source_inspection_report(
    report: dict[str, Any],
    reports_dir: str | Path = DEFAULT_EGP_REPORTS_DIR,
) -> Path:
    path = Path(reports_dir) / SOURCE_INSPECTION_REPORT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(report, indent=2, ensure_ascii=False, sort_keys=True),
        encoding="utf-8",
    )
    return path


def _inspect_rows(
    rows: list[tuple[Any, ...]],
    headers: list[str],
    field_columns: dict[str, int],
    category: EGPCategoryConfig,
    header_row_number: int,
) -> dict[str, Any]:
    super_categories: Counter[str] = Counter()
    sub_categories: Counter[str] = Counter()
    cefr_distribution: Counter[str] = Counter()
    duplicate_groups: dict[str, list[int]] = defaultdict(list)
    malformed_rows: list[dict[str, Any]] = []
    category_mismatch_rows: list[dict[str, Any]] = []
    empty_rows = 0
    data_row_count = 0

    for offset, row in enumerate(rows, start=1):
        row_number = header_row_number + offset
        if _is_empty_row(row):
            empty_rows += 1
            continue

        data_row_count += 1
        super_category = _field_value(row, field_columns, "super_category")
        sub_category = _field_value(row, field_columns, "sub_category")
        cefr_level = (_field_value(row, field_columns, "cefr_level") or "").upper()
        can_do_statement = _field_value(row, field_columns, "can_do_statement")

        if super_category:
            super_categories[super_category] += 1
        if sub_category:
            sub_categories[sub_category] += 1
        if cefr_level:
            cefr_distribution[cefr_level] += 1

        row_issues = []
        category_issues = []
        if cefr_level not in CEFR_LEVELS:
            row_issues.append("invalid_cefr_level")
        if not can_do_statement:
            row_issues.append("empty_can_do_statement")
        if category.expected_super_category and not _source_label_matches(
            super_category,
            category.expected_super_category,
        ):
            category_issues.append("super_category_mismatch")
        if category.expected_sub_category and not _source_label_matches(
            sub_category,
            category.expected_sub_category,
        ):
            category_issues.append("sub_category_mismatch")
        if row_issues:
            malformed_rows.append(
                {
                    "source_row_number": row_number,
                    "issues": row_issues,
                },
            )
        if category_issues:
            category_mismatch_rows.append(
                {
                    "source_row_number": row_number,
                    "issues": category_issues,
                },
            )

        duplicate_groups[_row_duplicate_key(headers, row)].append(row_number)

    duplicate_row_count = sum(
        max(len(row_numbers) - 1, 0)
        for row_numbers in duplicate_groups.values()
    )
    return {
        "data_row_count": data_row_count,
        "super_category_values": dict(sorted(super_categories.items())),
        "sub_category_values": dict(sorted(sub_categories.items())),
        "cefr_distribution": dict(sorted(cefr_distribution.items())),
        "duplicate_row_count": duplicate_row_count,
        "empty_rows": empty_rows,
        "malformed_rows": malformed_rows,
        "category_mismatch_rows": category_mismatch_rows,
    }


def _schema_issues(headers: list[str], field_columns: dict[str, int]) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    required_fields = {"cefr_level", "can_do_statement"}
    for field_name in sorted(required_fields - field_columns.keys()):
        issues.append(
            {
                "code": f"missing_{field_name}",
                "severity": "error",
            },
        )

    known_aliases = set(CAN_DO_ALIASES)
    for aliases in FIELD_ALIASES.values():
        known_aliases.update(aliases)
    unexpected = [
        header
        for header in headers
        if normalize_column_name(header) not in known_aliases
    ]
    if unexpected:
        issues.append(
            {
                "code": "unexpected_columns",
                "severity": "warning",
                "message": ", ".join(unexpected),
            },
        )
    return issues


def _inspection_totals(files: list[dict[str, Any]]) -> dict[str, int]:
    return {
        "files_configured": len(files),
        "files_found": sum(1 for file in files if file["exists"]),
        "files_with_errors": sum(
            1
            for file in files
            if any(issue["severity"] == "error" for issue in file["issues"])
            or file["malformed_rows"]
        ),
        "data_rows": sum(int(file["data_row_count"]) for file in files),
        "malformed_rows": sum(len(file["malformed_rows"]) for file in files),
        "category_mismatch_rows": sum(
            len(file["category_mismatch_rows"])
            for file in files
        ),
        "duplicate_rows": sum(int(file["duplicate_row_count"]) for file in files),
        "empty_rows": sum(int(file["empty_rows"]) for file in files),
    }


def _field_value(
    row: tuple[Any, ...],
    field_columns: dict[str, int],
    field_name: str,
) -> str | None:
    index = field_columns.get(field_name)
    if index is None or index >= len(row):
        return None
    return normalize_empty(row[index])


def _row_duplicate_key(headers: list[str], row: tuple[Any, ...]) -> str:
    values = [
        _field_value(row, {header: index for index, header in enumerate(headers)}, header)
        for header in headers
    ]
    raw_key = "\x1f".join(_normalize_for_key(value) for value in values)
    return hashlib.sha256(raw_key.encode("utf-8")).hexdigest()


def _normalize_for_key(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().casefold().split())


def _source_label_matches(actual: str | None, expected: str | None) -> bool:
    return _normalize_for_key(actual) == _normalize_for_key(expected)


def _is_empty_row(row: tuple[Any, ...]) -> bool:
    return all(normalize_empty(value) is None for value in row)
