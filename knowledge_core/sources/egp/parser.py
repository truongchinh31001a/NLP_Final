from __future__ import annotations

import hashlib
import json
import logging
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from knowledge_core.sources.egp.models import EGPCategoryConfig, EGPConfig, RawEGPRecord
from knowledge_core.sources.egp.normalizer import (
    normalize_empty,
    normalize_record,
    parse_feature_label,
)
from knowledge_core.sources.egp.paths import DEFAULT_RAW_GRAMMAR_DIR, raw_file_path


logger = logging.getLogger(__name__)


CAN_DO_ALIASES = {
    "candostatement",
    "candodescriptor",
    "cando",
    "statement",
    "descriptor",
}
FIELD_ALIASES = {
    "source_record_id": {"id", "refid", "recordid", "sourceid", "sourcerecordid"},
    "cefr_level": {"level", "cefr", "cefrlevel", "cefrtext"},
    "super_category": {"supercategory", "supercategorytext"},
    "sub_category": {"subcategory", "subcategorytext"},
    "can_do_statement": CAN_DO_ALIASES,
    "example": {"example", "examples", "learnerexamples"},
    "details": {"comments", "comment", "details", "detail", "lexicalrange"},
    "feature": {"guideword", "feature", "featurelabel", "guidewordtext"},
}
HEADER_SCAN_ROWS = 20


class EGPParseError(ValueError):
    """Raised when an EGP XLSX workbook cannot be parsed safely."""


@dataclass(slots=True)
class ParsedEGPFile:
    category_id: str
    source_file: Path
    records: list[RawEGPRecord]
    empty: bool = False


@dataclass(slots=True)
class ParsedEGPDataset:
    files: list[ParsedEGPFile]

    @property
    def records(self) -> list[RawEGPRecord]:
        return [record for parsed_file in self.files for record in parsed_file.records]

    @property
    def empty_files(self) -> list[Path]:
        return [parsed_file.source_file for parsed_file in self.files if parsed_file.empty]


def normalize_column_name(value: Any) -> str:
    text = normalize_empty(value)
    if text is None:
        return ""
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def parse_workbook(
    source_file: str | Path,
    category: EGPCategoryConfig,
    *,
    source_name: str = "english_grammar_profile",
    retrieved_at: datetime | None = None,
    source_url: str | None = None,
) -> list[RawEGPRecord]:
    path = Path(source_file)
    if not path.exists():
        raise EGPParseError(f"EGP XLSX file does not exist: {path}")
    if path.stat().st_size == 0:
        raise EGPParseError(f"EGP XLSX file is empty: {path}")

    metadata = _load_source_metadata(path)
    record_retrieved_at = retrieved_at or _metadata_datetime(metadata) or _file_mtime(path)
    record_source_url = source_url or _metadata_source_url(metadata)

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise EGPParseError(f"Malformed EGP XLSX file {path}: {exc}") from exc

    try:
        worksheet = workbook.active
        header_row_number, headers = _find_header_row(worksheet.iter_rows(values_only=True))
        field_columns = _resolve_field_columns(headers)
        if "can_do_statement" not in field_columns:
            raise EGPParseError(
                f"EGP XLSX file {path} is missing a Can-DoStatement column",
            )

        records: list[RawEGPRecord] = []
        for excel_row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row_number + 1, values_only=True),
            start=header_row_number + 1,
        ):
            if _is_empty_row(row):
                continue
            raw_payload = _raw_payload(headers, row)
            record = _record_from_row(
                category=category,
                source_name=source_name,
                source_file=path,
                source_url=record_source_url,
                retrieved_at=record_retrieved_at,
                source_row_number=excel_row_number,
                field_columns=field_columns,
                row=row,
                raw_payload=raw_payload,
            )
            normalized_record = normalize_record(record)
            records.append(
                normalized_record.model_copy(
                    update={
                        "source_record_id": deterministic_source_record_id(
                            normalized_record,
                        ),
                    },
                ),
            )
        return records
    finally:
        workbook.close()


def parse_configured_workbooks(
    config: EGPConfig,
    raw_dir: str | Path = DEFAULT_RAW_GRAMMAR_DIR,
    *,
    category_id: str | None = None,
) -> ParsedEGPDataset:
    selected_categories = config.select_categories(category_id)
    raw_path = Path(raw_dir)
    parsed_files: list[ParsedEGPFile] = []

    for category in selected_categories:
        source_file = raw_file_path(raw_path, category)
        records = parse_workbook(
            source_file,
            category,
            source_name=config.source.name,
        )
        parsed_files.append(
            ParsedEGPFile(
                category_id=category.id,
                source_file=source_file,
                records=records,
                empty=len(records) == 0,
            ),
        )
        if not records:
            logger.warning("Parsed EGP workbook has no records: %s", source_file)

    return ParsedEGPDataset(parsed_files)


def _find_header_row(rows: Iterable[tuple[Any, ...]]) -> tuple[int, list[str]]:
    first_non_empty: tuple[int, tuple[Any, ...]] | None = None
    for row_number, row in enumerate(rows, start=1):
        if row_number > HEADER_SCAN_ROWS:
            break
        if _is_empty_row(row):
            continue
        if first_non_empty is None:
            first_non_empty = (row_number, row)
        normalized = {normalize_column_name(value) for value in row}
        if normalized & CAN_DO_ALIASES:
            return row_number, _headers_from_row(row)

    if first_non_empty is not None:
        row_number, row = first_non_empty
        return row_number, _headers_from_row(row)
    raise EGPParseError("EGP XLSX workbook contains no header row")


def _headers_from_row(row: tuple[Any, ...]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(row, start=1):
        text = normalize_empty(value) or f"column_{index}"
        count = seen.get(text, 0)
        seen[text] = count + 1
        headers.append(text if count == 0 else f"{text}_{count + 1}")
    return headers


def _resolve_field_columns(headers: list[str]) -> dict[str, int]:
    normalized_headers = [normalize_column_name(header) for header in headers]
    fields: dict[str, int] = {}
    for index, normalized_header in enumerate(normalized_headers):
        for field_name, aliases in FIELD_ALIASES.items():
            if field_name not in fields and normalized_header in aliases:
                fields[field_name] = index
    return fields


def _record_from_row(
    *,
    category: EGPCategoryConfig,
    source_name: str,
    source_file: Path,
    source_url: str | None,
    retrieved_at: datetime,
    source_row_number: int,
    field_columns: dict[str, int],
    row: tuple[Any, ...],
    raw_payload: dict[str, Any],
) -> RawEGPRecord:
    def get(field_name: str) -> str | None:
        index = field_columns.get(field_name)
        if index is None or index >= len(row):
            return None
        return _cell_to_text(row[index])

    feature_source = get("feature") or get("can_do_statement")
    feature_type, feature_name = parse_feature_label(feature_source)
    return RawEGPRecord(
        source=source_name,
        source_record_id=None,
        category_id=category.id,
        super_category=get("super_category"),
        sub_category=get("sub_category"),
        cefr_level=get("cefr_level") or "",
        feature_type=feature_type,
        feature_name=feature_name,
        can_do_statement=get("can_do_statement") or "",
        example=get("example"),
        details=get("details"),
        source_row_number=source_row_number,
        source_file=str(source_file),
        source_url=source_url,
        retrieved_at=retrieved_at,
        canonical_parent_hint=category.canonical_parent_hint,
        raw_payload=raw_payload,
    )


def deterministic_source_record_id(record: RawEGPRecord) -> str:
    parts = [
        record.source,
        Path(record.source_file).as_posix(),
        str(record.source_row_number or ""),
        record.category_id,
        record.super_category,
        record.sub_category,
        record.cefr_level,
        record.can_do_statement,
    ]
    raw_key = "\x1f".join(_record_id_part(part) for part in parts)
    return "egp_" + hashlib.sha256(raw_key.encode("utf-8")).hexdigest()


def _record_id_part(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def _raw_payload(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    payload: dict[str, Any] = {}
    for index, header in enumerate(headers):
        value = row[index] if index < len(row) else None
        payload[header] = _json_safe_cell(value)
    return payload


def _cell_to_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    return normalize_empty(value)


def _json_safe_cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def _is_empty_row(row: tuple[Any, ...]) -> bool:
    return all(normalize_empty(value) is None for value in row)


def _metadata_path(source_file: Path) -> Path:
    return source_file.with_suffix(".metadata.json")


def _load_source_metadata(source_file: Path) -> dict[str, Any] | None:
    metadata_path = _metadata_path(source_file)
    if not metadata_path.exists():
        return None
    try:
        payload = json.loads(metadata_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise EGPParseError(f"Invalid EGP metadata JSON {metadata_path}: {exc}") from exc
    if not isinstance(payload, dict):
        raise EGPParseError(f"EGP metadata must be an object: {metadata_path}")
    return payload


def _metadata_datetime(metadata: dict[str, Any] | None) -> datetime | None:
    if not metadata:
        return None
    value = metadata.get("retrieved_at")
    if not isinstance(value, str) or not value.strip():
        return None
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=timezone.utc)
    return parsed


def _metadata_source_url(metadata: dict[str, Any] | None) -> str | None:
    if not metadata:
        return None
    value = metadata.get("source_url")
    return normalize_empty(value)


def _file_mtime(source_file: Path) -> datetime:
    return datetime.fromtimestamp(source_file.stat().st_mtime, tz=timezone.utc)
